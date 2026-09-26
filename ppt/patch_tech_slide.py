r"""Swap the TECHNICAL APPROACH slide's two old flowchart pictures for the single
winner-style architecture diagram rendered by ppt/diagram/render.py.

Keeps the official template furniture untouched: title placeholder, banner image,
team-name oval, footer bar, footer text and slide-number placeholder, plus the two
mandated instruction pointers ("Technologies to be used" / "Methodology and process
for implementation") in the body text box.

Run:  .venv\Scripts\python.exe ppt\patch_tech_slide.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent.parent
PPT_DIR = ROOT / "ppt"
DECK = PPT_DIR / "final" / "SIH26120_Idea_Presentation_STRICT.pptx"
DIAGRAM = PPT_DIR / "diagram" / "tech_architecture.png"

SLIDE_IDX = 2  # TECHNICAL APPROACH

# --- content-area budget (inches, 13.333 x 7.5 canvas) ----------------------- #
POINTER_TOP, POINTER_H = 1.24, 0.84   # body text box holding the two pointers
DIA_TOP = 2.16                        # below the pointers
DIA_BOTTOM_LIMIT = 6.80               # footer bar starts at 6.95
SLIDE_W = 13.3333

# template furniture that the picture must never touch
FURNITURE = {
    "Title 1",
    "Oval 10",
    "Rectangle 9",
    "Picture 11",
    "Footer Placeholder 6",
    "Slide Number Placeholder 5",
    "TextBox 8",
}


def rect(shape):
    return (
        shape.left / 914400,
        shape.top / 914400,
        (shape.left + shape.width) / 914400,
        (shape.top + shape.height) / 914400,
    )


def overlaps(a, b):
    return not (a[2] <= b[0] or b[2] <= a[0] or a[3] <= b[1] or b[3] <= a[1])


def main() -> None:
    if not DIAGRAM.exists():
        sys.exit(f"missing {DIAGRAM} — run ppt/diagram/render.py first")

    prs = Presentation(str(DECK))
    slide = list(prs.slides)[SLIDE_IDX]

    title = slide.shapes.title.text_frame.text.strip()
    if title != "TECHNICAL APPROACH":
        sys.exit(f"slide {SLIDE_IDX} is {title!r}, not TECHNICAL APPROACH")

    # 1. tighten the pointer text box so the diagram gets maximum height, and make
    #    the "Technologies to be used" pointer match the diagram's STACK strip
    body = next(s for s in slide.shapes if s.name == "TextBox 8")
    body.top, body.height = Inches(POINTER_TOP), Inches(POINTER_H)
    tech_line = (
        "Technologies to be used: Python 3.13 · NumPy / SciPy · pandas · "
        "scikit-learn · XGBoost · scikit-optimize · FastAPI · Plotly "
        "· pytest"
    )
    for para in body.text_frame.paragraphs:
        if para.text.startswith("Technologies to be used"):
            para.runs[0].text = tech_line
            for r in para.runs[1:]:
                r._r.getparent().remove(r._r)
            break
    else:
        sys.exit("could not find the 'Technologies to be used' pointer line")

    # 2. drop the old system-flow + tech-stack pictures (keep the SIH banner)
    removed = []
    for shp in list(slide.shapes):
        if shp.shape_type == 13 and shp.name not in FURNITURE:  # PICTURE
            removed.append(shp.name)
            shp._element.getparent().remove(shp._element)
    print("removed pictures:", removed or "(none)")

    # 3. add the new diagram, height-fitted and horizontally centred
    im = Image.open(DIAGRAM)
    ar = im.width / im.height
    h = DIA_BOTTOM_LIMIT - DIA_TOP
    w = h * ar
    if w > 12.70:                       # would never happen at 1650x780, but be safe
        w, h = 12.70, 12.70 / ar
    left = (SLIDE_W - w) / 2
    pic = slide.shapes.add_picture(
        str(DIAGRAM), Inches(left), Inches(DIA_TOP), Inches(w), Inches(h)
    )
    pic.name = "TechArchitecture"
    print(f"placed {DIAGRAM.name}: {w:.2f}in x {h:.2f}in at ({left:.2f}, {DIA_TOP:.2f})")

    # 4. verify: no overlap with any template furniture, nothing off-canvas
    pr = rect(pic)
    problems = []
    if pr[0] < 0 or pr[1] < 0 or pr[2] > SLIDE_W or pr[3] > 7.5:
        problems.append(f"picture off-canvas: {pr}")
    for shp in slide.shapes:
        if shp._element is pic._element:
            continue
        r = rect(shp)
        if overlaps(pr, r):
            problems.append(f"picture overlaps {shp.name!r} {r}")
    body_r = rect(body)
    if body_r[3] > DIA_TOP:
        problems.append(f"pointer box bottom {body_r[3]:.2f} below diagram top {DIA_TOP}")

    for shp in slide.shapes:
        print(f"  {shp.name:<28} L={rect(shp)[0]:6.2f} T={rect(shp)[1]:6.2f} "
              f"R={rect(shp)[2]:6.2f} B={rect(shp)[3]:6.2f}")

    if problems:
        for p in problems:
            print("!! " + p)
        sys.exit("layout verification failed")
    print("layout OK — no overlap with template furniture, inside 13.333 x 7.5")

    prs.save(str(DECK))
    print(f"saved {DECK.name} ({DECK.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
