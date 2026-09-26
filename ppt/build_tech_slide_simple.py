r"""Slide 3 (TECHNICAL APPROACH) -- the SIMPLE, icon-based variant.

The detailed native wiring diagram (build_tech_slide_native.py) stays the slide in the
submission deck. This script builds its non-technical sibling: one left-to-right story
of five big icon stations (the well -> digital twin -> smart search -> dashboard ->
operator), a dashed "records come back" loop, and a strip of what the operator gets and
what it leads to. Same template furniture, same palette (navy / orange / green /
purple), same rounded panels with coloured header bands, same triangle arrowheads.

Output (a PREVIEW -- the STRICT submission deck is not touched):
    ppt/final/SIH26120_Slide3_SIMPLE_preview.pptx   (1 slide: the template's slide 3)
    ppt/final/SIH26120_Slide3_SIMPLE_preview.pdf
    ppt/diagram/slide3_simple.png                    (1920 x 1080)

How: open the STRICT deck, keep only slide 3 (title placeholder, SIH banner, team oval,
footer bar, footer text, slide number and the template's pointer text box are kept as
they are), delete the detailed diagram's `TA_*` shapes, draw this one as `TS_*` shapes.
Every word is a native text frame (copyable in the PDF); only the icons are pictures,
rasterised from ppt/assets/svg/icons/ through headless Edge (black render -> alpha mask
-> recoloured to the station colour, so they sit cleanly on the tinted discs).

Gates (nothing is saved if one trips) -- the native builder's five gates, imported from
build_tech_slide_native.py (canvas band, template furniture, node overlap, wiring, text
fit), plus the simple-slide rules:
  * every sentence >= 11 pt, station titles 14-16 pt, station icons >= 0.9 in;
  * body text (5 sentences + loop caption) <= 60 words, each sentence <= 12 words;
  * no acronyms / code names (no all-caps token, no '_', '/', '.py');
  * no orphan word on the last line of any wrapped sentence;
  * pictures and discs never overlap text, furniture or each other; no wire crosses
    a station panel it does not start or end at.
Then PowerPoint COM exports PDF + PNG and pypdf must find every station title,
sentence and chip label on the page as text.

Run:
    .venv\Scripts\python.exe ppt\build_tech_slide_simple.py              # preview deck + PDF + PNG
    .venv\Scripts\python.exe ppt\build_tech_slide_simple.py --no-export  # gates + .pptx only
    .venv\Scripts\python.exe ppt\build_tech_slide_simple.py --into-strict
        # SWAP: replace slide 3 of the STRICT submission deck with this simple
        # version (backs the deck up to ppt/archive/ first). Not run by default --
        # the detailed slide is the one being submitted. To go back, copy the
        # backup over the deck (build_tech_slide_native.py refuses to run on a
        # slide 3 that carries TS_* shapes, by design).
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE

ROOT = Path(__file__).resolve().parent.parent
PPT_DIR = ROOT / "ppt"
sys.path.insert(0, str(PPT_DIR))
import build_tech_slide_native as N  # noqa: E402  (gates, text fit, drawing helpers)
from build_tech_slide_native import (  # noqa: E402
    A, Box, E, EMU, Para, Wire, C_IN, C_ML, C_PHY, C_SVC, C_TEXT, PANEL_FILL, TINT,
    FURNITURE, SLIDE_IDX, inside, overlaps,
)

STRICT = PPT_DIR / "final" / "SIH26120_Idea_Presentation_STRICT.pptx"
STRICT_PDF = STRICT.with_suffix(".pdf")
OUT = PPT_DIR / "final" / "SIH26120_Slide3_SIMPLE_preview.pptx"
OUT_PDF = OUT.with_suffix(".pdf")
PNG = PPT_DIR / "diagram" / "slide3_simple.png"
SWAP_BACKUP = PPT_DIR / "archive" / "SIH26120_Idea_Presentation_STRICT_pre-simple-slide3.pptx"
ICON_DIR = PPT_DIR / "assets" / "svg" / "icons"
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"

PREFIX = "TS_"
C_ARROW = "595959"
C_LOOP = "404040"
C_NOTE = "595959"

# --------------------------------------------------------------------------- #
# icons: library SVGs (and three compositions of library parts)               #
# --------------------------------------------------------------------------- #
SVG_HEAD = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 48 48" fill="none" '
            'stroke="currentColor" stroke-width="{sw}" stroke-linecap="round" '
            'stroke-linejoin="round">')


def _inner(name: str) -> str:
    """The drawing elements of ppt/assets/svg/icons/ic_<name>.svg (wrapper stripped)."""
    txt = (ICON_DIR / f"ic_{name}.svg").read_text(encoding="utf-8")
    return re.sub(r"</svg>\s*$", "", re.sub(r"^\s*<svg[^>]*>", "", txt)).strip()


def _lib(name: str, sw: float = 2.0) -> str:
    return SVG_HEAD.format(sw=sw) + _inner(name) + "</svg>"


def _operator() -> str:
    # ic_team's centre figure on its own, lowered a little, wearing a hard hat
    return (SVG_HEAD.format(sw=2.2)
            + '<circle cx="24" cy="18.5" r="5.5"/>'
            + '<path d="M12 42 C12 32.5 17 28 24 28 C31 28 36 32.5 36 42"/>'
            + '<path d="M16.5 13 A7.5 7.5 0 0 1 31.5 13"/>'
            + '<line x1="14.5" y1="13" x2="33.5" y2="13"/>'
            + '<line x1="24" y1="5.5" x2="24" y2="9"/>'
            + "</svg>")


def _twin() -> str:
    # ic_dashboard's screen frame (plus a stand) with ic_pumpjack drawn inside it:
    # "a computer copy of the well"
    return (SVG_HEAD.format(sw=2.2)
            + '<rect x="3" y="6" width="42" height="30" rx="3"/>'
            + '<line x1="24" y1="36" x2="24" y2="42"/>'
            + '<line x1="16" y1="42" x2="32" y2="42"/>'
            + '<g transform="translate(9.9 5.6) scale(0.6)" stroke-width="3.3">'
            + _inner("pumpjack") + "</g></svg>")


def _search() -> str:
    # a magnifying glass whose lens holds ic_neural-net: "smart search"
    return (SVG_HEAD.format(sw=2.2)
            + '<circle cx="20" cy="20" r="15"/>'
            + '<line x1="31" y1="31" x2="43" y2="43" stroke-width="4"/>'
            + '<g transform="translate(5.12 5.12) scale(0.62)" stroke-width="2.6">'
            + _inner("neural-net") + "</g></svg>")


ICONS = {  # key -> (svg markup, where it comes from)
    "well": (_lib("pumpjack", 2.2), "ic_pumpjack.svg"),
    "twin": (_twin(), "ic_dashboard.svg frame + ic_pumpjack.svg (composed)"),
    "search": (_search(), "magnifier + ic_neural-net.svg (composed)"),
    "dash": (_lib("dashboard", 2.2), "ic_dashboard.svg"),
    "operator": (_operator(), "ic_team.svg centre figure + hard hat (composed)"),
    "loop": (_lib("calendar-cycle", 2.4), "ic_calendar-cycle.svg"),
    "steam": (_lib("steam", 2.6), "ic_steam.svg"),
    "pump": (_lib("pumpjack", 2.6), "ic_pumpjack.svg"),
    "alert": (_lib("alert-triangle", 2.6), "ic_alert-triangle.svg"),
    "flame": (_lib("flame", 2.6), "ic_flame.svg"),
    "shield": (_lib("shield-check", 2.6), "ic_shield-check.svg"),
    "oil": (_lib("oil-drop", 2.6), "ic_oil-drop.svg"),
}


def rasterise(key: str, color: str, work: Path, px: int = 480) -> Path:
    """SVG -> PNG with headless Edge (black on white), then the grey level becomes the
    alpha channel and the colour is set to `color` -- anti-aliasing is kept and the
    icon has a transparent background."""
    svg, _src = ICONS[key]
    svg_path = work / f"{key}.svg"
    svg_path.write_text(svg.replace("currentColor", "#000000"), encoding="utf-8")
    html = work / f"{key}.html"
    html.write_text("<style>html,body{margin:0;padding:0;background:#fff;overflow:hidden}"
                    f"img{{display:block;width:{px}px;height:{px}px}}</style>"
                    f'<img src="file:///{svg_path.as_posix()}">', encoding="utf-8")
    raw = work / f"{key}_raw.png"
    if raw.exists():
        raw.unlink()
    subprocess.run([EDGE, "--headless", "--disable-gpu", "--hide-scrollbars",
                    "--force-device-scale-factor=1", f"--screenshot={raw}",
                    f"--window-size={px},{px}", f"file:///{html.as_posix()}"],
                   check=False, capture_output=True)
    if not raw.exists():
        raise RuntimeError(f"Edge failed to render icon {key}")
    grey = Image.open(raw).convert("L").crop((0, 0, px, px))
    alpha = grey.point(lambda v: 255 - v)
    rgb = tuple(int(color[i:i + 2], 16) for i in (0, 2, 4))
    out = Image.new("RGBA", grey.size, rgb + (0,))
    out.putalpha(alpha)
    if out.getbbox() is None:
        raise RuntimeError(f"icon {key} rendered empty")
    png = work / f"{key}_{color}.png"
    out.save(png)
    return png


# --------------------------------------------------------------------------- #
# content                                                                      #
# --------------------------------------------------------------------------- #
STATIONS = [  # key, colour, title, sentence as hand-balanced lines (one paragraph each)
    ("well", C_IN, "The well",
     ("Oil India's heavy-oil well:", "steam goes in,", "oil comes out.")),
    ("twin", C_PHY, "Digital twin",
     ("A computer copy", "of the well —", "heat, thick oil, the pump.")),
    ("search", C_ML, "Smart search",
     ("Tries 53,000 settings", "in 36 seconds,", "keeps the safe ones.")),
    ("dash", C_SVC, "Dashboard",
     ("Shows the best plan", "in plain words,", "English and Hindi.")),
    ("operator", C_IN, "Operator",
     ("Decides.", "Nothing is sent", "to the well automatically.")),
]
LOOP_TEXT = "Real cycle records come back → the twin learns"
GETS = [("steam", ("How much steam", "& pressure")), ("pump", ("Pump speed", "& stroke")),
        ("alert", ("When to stop before", "the rods jam"))]
LEADS = [("flame", ("Less diesel", "per barrel")), ("shield", ("Safer pump",)),
         ("oil", ("More oil", "per cycle"))]
GETS_TITLE, LEADS_TITLE = "What the operator gets", "What it leads to"
TRUST = ("Built on published physics · checked against real field data · "
         "236 automated tests")


def joined(lines) -> str:
    return " ".join(lines)


# --------------------------------------------------------------------------- #
# layout                                                                       #
# --------------------------------------------------------------------------- #
X0, X1 = 0.32, 13.11
CARD_W = 2.15
GAP = (X1 - X0 - 5 * CARD_W) / 4
CARD_T = 2.22
HDR_H = 0.42
DISC_D = 1.12
ICON_D = 0.94
SENT_H = 0.64
CARD_H = HDR_H + 0.10 + DISC_D + 0.08 + SENT_H + 0.05
TITLE_PT, SENT_PT, CHIP_PT = 15.0, 12.0, 12.0

LOOP_Y = CARD_T + CARD_H + 0.26
STRIP_T = LOOP_Y + 0.46
STRIP_HDR = 0.32
CHIP_H = 0.62
STRIP_H = STRIP_HDR + 0.10 + CHIP_H + 0.10
STRIP_GAP = 0.40

PICS: list[tuple] = []     # (id, key, colour, x, y, size, container-id)
DISCS: list[tuple] = []    # (id, colour, x, y, d)
BIG: list[Wire] = []       # the fat station arrows (drawn with large heads)


def card_x(i):
    return X0 + i * (CARD_W + GAP)


def layout():
    N.BOXES.clear()
    N.WIRES.clear()
    N.DOTS.clear()
    # ---- five stations ---------------------------------------------------------
    for i, (key, col, title, sentence) in enumerate(STATIONS):
        x = card_x(i)
        pid = f"S{i + 1}"
        N.add(Box(pid, x, CARD_T, CARD_W, CARD_H, "panel", col, fill=PANEL_FILL[col],
                  line=col, line_w=1.5, radius=0.12))
        N.add(Box(pid + "_hdr", x, CARD_T, CARD_W, HDR_H, "header", col,
                  [Para(f"{i + 1} · {title}", TITLE_PT, True, color="FFFFFF",
                        align="ctr")],
                  fill=col, line=None, radius=0.12, anchor="ctr",
                  ins=(0.06, 0.0, 0.06, 0.0)))
        dy = CARD_T + HDR_H + 0.10
        dx = x + (CARD_W - DISC_D) / 2
        DISCS.append((pid + "_disc", col, dx, dy, DISC_D))
        PICS.append((pid + "_icon", key, col, dx + (DISC_D - ICON_D) / 2,
                     dy + (DISC_D - ICON_D) / 2, ICON_D, pid))
        N.add(Box(pid + "_txt", x + 0.08, dy + DISC_D + 0.08, CARD_W - 0.16, SENT_H,
                  "label", col, [Para(t, SENT_PT, color=C_TEXT, align="ctr") for t in sentence],
                  fill=None, line=None, anchor="t", ins=(0.0, 0.0, 0.0, 0.0)))
    # ---- big arrows between stations ------------------------------------------
    ay = CARD_T + HDR_H + 0.10 + DISC_D / 2
    for i in range(4):
        a, b = f"S{i + 1}", f"S{i + 2}"
        w = Wire([(card_x(i) + CARD_W + 0.05, ay), (card_x(i + 1) - 0.05, ay)], a, b,
                 kind="flow")
        N.WIRES.append(w)
        BIG.append(w)
    # ---- the loop: records come back, the twin learns --------------------------
    op, tw = "S5", "S2"
    ox, tx = card_x(4) + CARD_W / 2, card_x(1) + CARD_W / 2
    bottom = CARD_T + CARD_H
    N.WIRES.append(Wire([(ox, bottom), (ox, LOOP_Y), (tx, LOOP_Y), (tx, bottom)], op, tw,
                        kind="feedback"))
    lw = 4.30
    lx = (ox + tx) / 2 - lw / 2 + 0.2
    N.add(Box("loop_txt", lx, LOOP_Y + 0.07, lw, 0.26, "label", C_LOOP,
              [Para(LOOP_TEXT, SENT_PT, True, color=C_LOOP, align="l")],
              fill=None, line=None, anchor="ctr", ins=(0.0, 0.0, 0.0, 0.0)))
    PICS.append(("loop_icon", "loop", C_LOOP, lx - 0.40, LOOP_Y + 0.03, 0.34, None))
    # ---- bottom strip: what the operator gets -> what it leads to --------------
    gets_w = 7.24                               # the "gets" labels are longer
    widths = (gets_w, X1 - X0 - STRIP_GAP - gets_w)
    for gi, (gid, col, title, items) in enumerate(
            [("G", C_IN, GETS_TITLE, GETS), ("L", C_ML, LEADS_TITLE, LEADS)]):
        sw = widths[gi]
        gx = X0 if gi == 0 else X1 - sw
        N.add(Box(gid, gx, STRIP_T, sw, STRIP_H, "panel", col, fill=PANEL_FILL[col],
                  line=col, line_w=1.25, radius=0.10))
        N.add(Box(gid + "_hdr", gx, STRIP_T, sw, STRIP_HDR, "header", col,
                  [Para(title, 13.0, True, color="FFFFFF")],
                  fill=col, line=None, radius=0.10, anchor="ctr",
                  ins=(0.12, 0.0, 0.06, 0.0)))
        pad, cg = 0.10, 0.10
        cw = (sw - 2 * pad - 2 * cg) / 3
        cy = STRIP_T + STRIP_HDR + 0.10
        for ci, (ikey, text) in enumerate(items):
            cx = gx + pad + ci * (cw + cg)
            cid = f"{gid}{ci + 1}"
            N.add(Box(cid, cx, cy, cw, CHIP_H, "chip", col,
                      [Para(t, CHIP_PT, True, color=C_TEXT) for t in text], fill="FFFFFF",
                      line=col,
                      line_w=0.75, panel=gid, anchor="ctr", radius=0.06,
                      ins=(0.54, 0.02, 0.05, 0.02)))
            PICS.append((cid + "_icon", ikey, col, cx + 0.08, cy + (CHIP_H - 0.42) / 2,
                         0.42, cid))
    # gets -> leads to
    ly = STRIP_T + STRIP_HDR + 0.10 + CHIP_H / 2
    w = Wire([(X0 + gets_w + 0.05, ly), (X0 + gets_w + STRIP_GAP - 0.05, ly)], "G", "L")
    N.WIRES.append(w)
    BIG.append(w)
    # ---- trust line ---------------------------------------------------------------
    ty = STRIP_T + STRIP_H + 0.08
    N.add(Box("trust", X1 - 7.2, ty, 7.2, 0.24, "label", C_NOTE,
              [Para(TRUST, 11.0, italic=True, color=C_NOTE, align="r")],
              fill=None, line=None, anchor="ctr", ins=(0.0, 0.0, 0.0, 0.0)))


# --------------------------------------------------------------------------- #
# extra gates for the simple slide                                             #
# --------------------------------------------------------------------------- #
def words(s: str) -> int:
    return len([t for t in re.split(r"\s+", s) if re.search(r"[A-Za-z0-9]", t)])


def simple_gates(furniture):
    problems = []
    texts = [b for b in N.BOXES.values() if b.paras]
    # fonts
    for b in texts:
        for p in b.paras:
            if p.size < 11.0:
                problems.append(f"[font] {b.id}: {p.size} pt < 11 pt")
    for i in range(5):
        sz = N.BOXES[f"S{i + 1}_hdr"].paras[0].size
        if not 14.0 <= sz <= 16.0:
            problems.append(f"[font] station title S{i + 1} is {sz} pt (want 14-16)")
    # words
    sentences = [joined(s) for *_, s in STATIONS]
    body = sentences + [LOOP_TEXT]
    n_body = sum(words(s) for s in body)
    if n_body > 60:
        problems.append(f"[words] body text is {n_body} words (> 60)")
    for s in sentences:
        if words(s) > 12:
            problems.append(f"[words] sentence > 12 words: {s!r}")
    # jargon
    all_text = [p.text for b in texts for p in b.paras]
    for t in all_text:
        caps = re.findall(r"\b[A-Z]{2,}\b", t)
        if caps or "_" in t or "/" in t or ".py" in t:
            problems.append(f"[jargon] {t!r} ({caps or 'code-like'})")
    # orphans: every hand-set line must stay ONE rendered line, and no line may be a
    # lone word unless it is a whole sentence ("Decides.") opening the text
    for b in texts:
        g = N.geom_inset(b)
        width_pt = (b.w - b.ins[0] - b.ins[2] - 2 * g) * 72 * N.FIT_SLACK
        for k, p in enumerate(b.paras):
            lines = N.wrap(p, width_pt) or []
            if len(b.paras) > 1 and len(lines) != 1:
                problems.append(f"[orphan] {b.id}: hand-set line {p.text!r} wraps to {lines}")
            for j, ln in enumerate(lines):
                if len(b.paras) + len(lines) > 2 and len(ln.split()) < 2 and not (
                        k == 0 and j == 0 and ln.endswith(".")):
                    problems.append(f"[orphan] {b.id}: lone word {ln!r} in {p.text!r}")
    # icons
    for pid, key, _c, x, y, d, cont in PICS:
        if cont and cont.startswith("S") and d < 0.9:
            problems.append(f"[icon] {pid} is {d:.2f} in (< 0.9 in)")
    # pictures + discs: band, furniture, text, each other
    rects = [(p[0], (p[3], p[4], p[3] + p[5], p[4] + p[5]), p[6]) for p in PICS]
    rects += [(d[0], (d[2], d[3], d[2] + d[4], d[3] + d[4]), None) for d in DISCS]
    for rid, r, cont in rects:
        if r[0] < N.BAND[0] or r[1] < N.BAND[1] or r[2] > N.BAND[2] or r[3] > N.BAND[3]:
            problems.append(f"[canvas] {rid} outside band")
        for fname, fr in furniture.items():
            if overlaps(r, fr):
                problems.append(f"[furniture] {rid} overlaps {fname}")
        for b in texts:
            tr = b.r
            if cont and b.id == cont:          # chip icon: must stay left of the text inset
                if r[2] > b.x + b.ins[0] - 0.02 or not inside(r, b.r):
                    problems.append(f"[icon] {rid} intrudes on {b.id}'s text")
                continue
            if overlaps(r, tr):
                problems.append(f"[icon] {rid} overlaps text box {b.id}")
    for i, (ia, ra, _) in enumerate(rects):
        for ib, rb, _ in rects[i + 1:]:
            nested = inside(ra, rb) or inside(rb, ra)
            if overlaps(ra, rb) and not nested:
                problems.append(f"[icon] {ia} overlaps {ib}")
    # wires must not cross a station / strip panel they do not belong to
    panels = [b for b in N.BOXES.values() if b.kind == "panel"]
    for w in N.WIRES:
        for p, q in zip(w.pts, w.pts[1:]):
            for pb in panels:
                if pb.id in (w.src, w.dst):
                    continue
                if N.seg_hits_rect(p, q, pb.r):
                    problems.append(f"[wiring] {w.src}->{w.dst} crosses panel {pb.id}")
    return problems, n_body


# --------------------------------------------------------------------------- #
# drawing                                                                      #
# --------------------------------------------------------------------------- #
def _big_ln(color, width_pt, dash=None):
    ln = etree.SubElement(etree.Element("x"), f"{A}ln", w=str(int(width_pt * 12700)))
    fill = etree.SubElement(ln, f"{A}solidFill")
    etree.SubElement(fill, f"{A}srgbClr", val=color)
    if dash:
        etree.SubElement(ln, f"{A}prstDash", val=dash)
    etree.SubElement(ln, f"{A}round")
    etree.SubElement(ln, f"{A}tailEnd", type="triangle", w="med", len="med")
    return ln


def draw(slide, work: Path):
    N.PREFIX = PREFIX
    boxes = list(N.BOXES.values())
    for b in boxes:
        if b.kind in ("panel", "header"):
            N.draw_box(slide, b)
    for did, col, x, y, d in DISCS:
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(x), E(y), E(d), E(d))
        shp.name = PREFIX + did
        N._strip_style(shp)
        N._set_fill(shp, "FFFFFF")
        N._set_ln(shp, N._ln_xml(col, 1.5))
    for i, w in enumerate(N.WIRES):
        (x1, y1) = w.pts[0]
        if w in BIG:
            (x2, y2) = w.pts[-1]
            shp = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2),
                                             E(y2))
            shp.name = f"{PREFIX}arrow{i:02d}_{w.src}_{w.dst}"
            N._strip_style(shp)
            N._set_ln(shp, _big_ln(C_ARROW, 4.0))
        else:
            fb = slide.shapes.build_freeform(E(x1), E(y1), scale=1.0)
            fb.add_line_segments([(E(x), E(y)) for x, y in w.pts[1:]], close=False)
            shp = fb.convert_to_shape()
            N._set_fill(shp, None)
            shp.name = f"{PREFIX}loop{i:02d}_{w.src}_{w.dst}"
            N._strip_style(shp)
            N._set_ln(shp, _big_ln(C_LOOP, 2.25, "dash"))
    for b in boxes:
        if b.kind not in ("panel", "header"):
            N.draw_box(slide, b)
    for pid, key, col, x, y, d, _cont in PICS:
        png = rasterise(key, col, work)
        pic = slide.shapes.add_picture(str(png), E(x), E(y), E(d), E(d))
        pic.name = PREFIX + pid
        pic._element.nvPicPr.cNvPr.set("descr", f"icon: {ICONS[key][1]}")


# --------------------------------------------------------------------------- #
# deck plumbing, export, QC                                                    #
# --------------------------------------------------------------------------- #
def open_slide3(keep_only: bool):
    prs = Presentation(str(STRICT))
    if keep_only:
        lst = prs.slides._sldIdLst
        for i, sld in reversed(list(enumerate(list(lst)))):
            if i != SLIDE_IDX:
                prs.part.drop_rel(sld.rId)
                lst.remove(sld)
        slide = prs.slides[0]
    else:
        slide = prs.slides[SLIDE_IDX]
    title = slide.shapes.title.text_frame.text.strip()
    if title != "TECHNICAL APPROACH":
        sys.exit(f"expected TECHNICAL APPROACH, got {title!r}")
    removed = 0
    for shp in list(slide.shapes):
        if shp.name.startswith(("TA_", "TS_", "TM_", "TF_")) or shp.name == N.OLD_PICTURE:
            shp._element.getparent().remove(shp._element)
            removed += 1
    xml = etree.tostring(slide._element).decode()
    for rId, rel in list(slide.part.rels.items()):
        if rel.reltype.endswith("/image") and f'"{rId}"' not in xml:
            slide.part.drop_rel(rId)
    furniture = {}
    for shp in slide.shapes:
        if shp.name in FURNITURE:
            furniture[shp.name] = (shp.left / EMU, shp.top / EMU,
                                   (shp.left + shp.width) / EMU,
                                   (shp.top + shp.height) / EMU)
    if FURNITURE - set(furniture):
        sys.exit(f"template furniture missing: {FURNITURE - set(furniture)}")
    extra = [s.name for s in slide.shapes if s.name not in FURNITURE]
    if extra:
        sys.exit(f"unexpected shapes left on slide 3: {extra}")
    print(f"opened {STRICT.name}: removed {removed} diagram shapes"
          f"{' , kept only slide 3' if keep_only else ''}")
    return prs, slide, furniture


def export_com(deck: Path, pdf: Path, slide_no: int):
    ps = f"""
$ErrorActionPreference = 'Stop'
$pp = New-Object -ComObject PowerPoint.Application
try {{
  $pres = $pp.Presentations.Open('{deck}', -1, 0, 0)
  $pres.SaveAs('{pdf}', 32)
  $pres.Slides.Item({slide_no}).Export('{PNG}', 'PNG', 1920, 1080)
  $pres.Close()
}} finally {{
  $pp.Quit()
}}
"""
    subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],
                   check=True)


def pdf_check(pdf: Path, page: int):
    from pypdf import PdfReader
    reader = PdfReader(str(pdf))
    raw = reader.pages[page].extract_text() or ""
    flat, tight = " ".join(raw.split()), "".join(raw.split())
    wanted = [f"{i + 1} · {t}" for i, (_k, _c, t, _s) in enumerate(STATIONS)]
    wanted += [joined(s) for *_, s in STATIONS] + [LOOP_TEXT, TRUST, GETS_TITLE,
                                                    LEADS_TITLE]
    wanted += [joined(t) for _k, t in GETS + LEADS]
    found = [w for w in wanted
             if " ".join(w.split()) in flat or "".join(w.split()) in tight]
    return len(reader.pages), wanted, found


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-export", action="store_true")
    ap.add_argument("--into-strict", action="store_true",
                    help="replace slide 3 of the STRICT submission deck (not the default)")
    args = ap.parse_args()

    into = args.into_strict
    prs, slide, furniture = open_slide3(keep_only=not into)
    layout()
    problems, fit_rows = N.run_gates(furniture)
    extra, n_body = simple_gates(furniture)
    problems += extra
    print("-- text fit (fullest 6) --")
    for fill, bid, detail in fit_rows[:6]:
        print(f"  {bid:<10} {fill * 100:5.1f}%  lines/para {detail}")
    if problems:
        print("\n!! GATES FAILED")
        for p in problems:
            print("  " + p)
        sys.exit(1)
    n_titles = sum(words(t) for *_, t, _s in STATIONS)
    print(f"GATES PASS: native gates (band, furniture {sorted(furniture)}, overlap, wiring, "
          f"text fit) + simple gates (fonts, words, jargon, orphans, icons)")
    print(f"words: body {n_body} (5 sentences + loop caption) · station titles {n_titles} "
          f"· trust line {words(TRUST)} · chips {sum(words(joined(t)) for _k, t in GETS + LEADS)}")

    work = Path(tempfile.mkdtemp(prefix="ts_icons_"))
    try:
        draw(slide, work)
    finally:
        shutil.rmtree(work, ignore_errors=True)

    if into:
        if not SWAP_BACKUP.exists():
            shutil.copy2(STRICT, SWAP_BACKUP)
            print(f"backed up STRICT deck -> {SWAP_BACKUP.relative_to(ROOT)}")
        deck, pdf, slide_no, pages = STRICT, STRICT_PDF, SLIDE_IDX + 1, 6
    else:
        deck, pdf, slide_no, pages = OUT, OUT_PDF, 1, 1
    prs.save(str(deck))
    print(f"saved {deck.relative_to(ROOT)} ({deck.stat().st_size / 1024:.0f} KB)")
    if args.no_export:
        return
    export_com(deck, pdf, slide_no)
    print(f"exported {pdf.relative_to(ROOT)} and {PNG.relative_to(ROOT)}")
    n_pages, wanted, found = pdf_check(pdf, slide_no - 1)
    print(f"PDF pages: {n_pages}; selectable text found: {len(found)}/{len(wanted)}")
    missing = [w for w in wanted if w not in found]
    if missing:
        print("  missing:", missing)
    if n_pages != pages or missing:
        sys.exit("PDF check failed")


if __name__ == "__main__":
    main()
