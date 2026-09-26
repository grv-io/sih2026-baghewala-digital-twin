"""
Build the EDITABLE variant of the visual deck.

    .venv\\Scripts\\python.exe ppt\\visual\\make_editable.py

For each slide it:
  1. copies slide<N>.html into editable/slide<N>.html with a measuring script
     appended. The script walks the DOM, records every text-bearing element's
     rect / font / colour / runs, then paints that text `transparent` — so the
     page keeps every card, chip, rule, chart and illustration but loses the copy.
  2. dumps the measured JSON out of the page with Edge `--dump-dom`.
  3. screenshots the text-free page to editable/bg<N>.png (2560x1440).
  4. assembles SIH26120_Idea_Presentation_VISUAL_EDITABLE.pptx — background
     picture per slide, plus a native PowerPoint text box for every measured
     element, at the matching position, size, font, weight and colour.

Text baked into the background on purpose: the navy left rail, and labels that
live inside an SVG illustration (axis ticks, leader labels, phase bands, bar
values). Those are geometry-locked artwork — edit them in the HTML and re-render.
"""

import json
import re
import shutil
import subprocess
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "editable"
PPT = HERE.parent
EDGE = Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe")
USER_DATA = HERE / ".edge-profile"
N = 7
EMU_PER_PX = 12192000 / 1280          # 9525 — 1280px canvas == 13.333in
PT_PER_PX = 0.75                       # 1280px == 960pt
PPTX_OUT = PPT / "archive" / "SIH26120_Idea_Presentation_VISUAL_EDITABLE.pptx"

# ---------------------------------------------------------------- measuring JS
MEASURE = r"""
<script>
(function () {
  var out = [], taken = [];
  function inTaken(el) {
    for (var i = 0; i < taken.length; i++) if (taken[i].contains(el)) return true;
    return false;
  }
  function runs(el) {
    var r = [];
    Array.prototype.forEach.call(el.childNodes, function (n) {
      var t = (n.textContent || '').replace(/\s+/g, ' ');
      if (!t.trim()) return;
      if (n.nodeType === 3) { r.push({ t: t, b: false, i: false, c: null }); }
      else if (n.nodeType === 1) {
        var cs = getComputedStyle(n);
        r.push({ t: t, b: parseInt(cs.fontWeight) >= 600,
                 i: cs.fontStyle === 'italic', c: cs.color });
      }
    });
    return r;
  }
  Array.prototype.forEach.call(document.querySelectorAll('.slide *'), function (el) {
    if (el.closest('.rail')) return;                 // rail stays baked in
    if (el.closest('svg')) return;                   // SVG labels stay baked in
    if (inTaken(el)) return;                         // already inside a captured box
    var direct = false;
    Array.prototype.forEach.call(el.childNodes, function (n) {
      if (n.nodeType === 3 && n.textContent.trim()) direct = true;
    });
    // an element whose children are purely inline runs (<b>/<em>/<small>) counts
    // as one text box. `querySelector('svg')` guards the case where the only
    // child is an <svg> — which also computes to display:inline, and would
    // otherwise swallow every chart label into one bogus box.
    if (!direct && el.children.length && !el.querySelector('svg')) {
      var allInline = Array.prototype.every.call(el.children, function (c) {
        return getComputedStyle(c).display === 'inline';
      });
      if (allInline && el.textContent.trim()) direct = true;
    }
    if (!direct) return;
    var r = el.getBoundingClientRect();
    if (r.width < 2 || r.height < 2) return;
    var cs = getComputedStyle(el);
    out.push({
      tag: el.tagName, cls: String(el.className || ''),
      x: r.left, y: r.top, w: r.width, h: r.height,
      fs: parseFloat(cs.fontSize), fw: parseInt(cs.fontWeight) || 400,
      color: cs.color, align: cs.textAlign,
      lh: cs.lineHeight, ls: cs.letterSpacing, tt: cs.textTransform,
      italic: cs.fontStyle === 'italic',
      mono: /consolas|courier/i.test(cs.fontFamily),
      runs: runs(el)
    });
    taken.push(el);
    // hide the element AND every descendant: inline runs such as <b>/<em>
    // carry their own CSS colour and would not inherit `transparent`.
    el.style.setProperty('color', 'transparent', 'important');
    Array.prototype.forEach.call(el.querySelectorAll('*'), function (c) {
      c.style.setProperty('color', 'transparent', 'important');
    });
  });
  var s = document.createElement('script');
  s.type = 'application/json'; s.id = '__meta';
  s.textContent = JSON.stringify(out);
  document.body.appendChild(s);
})();
</script>
"""

META_RE = re.compile(
    r'<script type="application/json" id="__meta">(.*?)</script>', re.S)


def edge(args, capture=False):
    cmd = [str(EDGE), "--headless=new", "--disable-gpu", "--no-sandbox",
           "--hide-scrollbars", "--force-color-profile=srgb",
           f"--user-data-dir={USER_DATA}", "--virtual-time-budget=4000"] + args
    return subprocess.run(cmd, capture_output=True, timeout=240,
                          text=capture, encoding="utf-8", errors="replace")


def css_rgb(s, default=(15, 37, 64)):
    m = re.match(r"rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)", s or "")
    if not m:
        return RGBColor(*default)
    return RGBColor(int(float(m.group(1))), int(float(m.group(2))), int(float(m.group(3))))


ALIGN = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT,
         "center": PP_ALIGN.CENTER, "justify": PP_ALIGN.JUSTIFY,
         "start": PP_ALIGN.LEFT, "end": PP_ALIGN.RIGHT}


def build():
    OUT.mkdir(exist_ok=True)
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    blank = prs.slide_layouts[6]
    stats = []

    for n in range(1, N + 1):
        src = HERE / f"slide{n}.html"
        page = OUT / f"slide{n}.html"
        html = src.read_text(encoding="utf-8")
        assert "</body>" in html
        page.write_text(html.replace("</body>", MEASURE + "\n</body>"),
                        encoding="utf-8")

        # --- measure ---
        r = edge(["--dump-dom", page.as_uri()], capture=True)
        m = META_RE.search(r.stdout or "")
        if not m:
            raise SystemExit(f"slide{n}: no __meta in dumped DOM")
        items = json.loads(m.group(1))

        # --- text-free background ---
        bg = OUT / f"bg{n}.png"
        if bg.exists():
            bg.unlink()
        edge([f"--screenshot={bg}", "--window-size=1280,720",
              "--force-device-scale-factor=2", page.as_uri()])
        if not bg.exists() or bg.stat().st_size < 50 * 1024:
            raise SystemExit(f"slide{n}: background render failed")
        with Image.open(bg) as im:
            assert im.size == (2560, 1440), f"slide{n}: {im.size}"

        # --- compose ---
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(bg), 0, 0,
                                 width=prs.slide_width, height=prs.slide_height)

        smallest = 99.0
        for it in items:
            # a little slack so PowerPoint's metrics don't force an early wrap,
            # clamped so no box ever runs off the 1280x720 canvas
            x0 = max(0.0, it["x"] - 2)
            y0 = max(0.0, it["y"] - 3)
            pad_w = max(8.0, it["w"] * 0.05)
            w = min(it["w"] + pad_w, 1272.0 - x0)
            h = min(it["h"] + 8, 716.0 - y0)
            left, top = Emu(int(x0 * EMU_PER_PX)), Emu(int(y0 * EMU_PER_PX))
            width, height = Emu(int(w * EMU_PER_PX)), Emu(int(h * EMU_PER_PX))
            box = slide.shapes.add_textbox(left, top, width, height)
            box.name = f"txt-{it['tag'].lower()}-{(it['cls'] or 'x').split()[0]}"
            tf = box.text_frame
            tf.word_wrap = True
            tf.auto_size = MSO_AUTO_SIZE.NONE
            tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

            p = tf.paragraphs[0]
            p.alignment = ALIGN.get(it["align"], PP_ALIGN.LEFT)
            lh = it.get("lh") or ""
            if lh.endswith("px"):
                p.line_spacing = Pt(round(float(lh[:-2]) * PT_PER_PX, 2))

            size_pt = round(it["fs"] * PT_PER_PX, 1)
            smallest = min(smallest, size_pt)
            base = css_rgb(it["color"])
            upper = it.get("tt") == "uppercase"
            ls = it.get("ls") or "normal"
            spc = None
            if ls.endswith("px"):
                spc = int(round(float(ls[:-2]) * PT_PER_PX * 100))

            for run in (it["runs"] or [{"t": "", "b": False, "i": False, "c": None}]):
                r_ = p.add_run()
                r_.text = run["t"].upper() if upper else run["t"]
                f = r_.font
                f.name = "Consolas" if it.get("mono") else "Segoe UI"
                f.size = Pt(size_pt)
                f.bold = bool(run["b"]) or it["fw"] >= 600
                f.italic = bool(run.get("i")) or bool(it.get("italic"))
                f.color.rgb = css_rgb(run["c"]) if run["c"] else base
                if spc:
                    f._rPr.set("spc", str(spc))

        stats.append((n, len(items), smallest))
        print(f"  slide{n}: {len(items):3d} text boxes, smallest {smallest:.1f} pt")

    prs.save(PPTX_OUT)
    total = sum(s[1] for s in stats)
    print(f"\nEDITABLE {PPTX_OUT}  {PPTX_OUT.stat().st_size/1024:.0f} KB  "
          f"({total} editable text boxes across {N} slides)")


if __name__ == "__main__":
    build()
