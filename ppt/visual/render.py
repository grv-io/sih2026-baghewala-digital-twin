"""
Baghewala Digital Twin — SIH26120 VISUAL deck renderer.

One command rebuilds everything:
    .venv\\Scripts\\python.exe ppt\\visual\\render.py

Pipeline:  slideN.html  --(headless Edge @2x)-->  slideN.png (2560x1440)
           slideN.png   --(python-pptx)-------->  SIH26120_Idea_Presentation_VISUAL.pptx
           slideN.png   --(Pillow)-------------->  SIH26120_Idea_Presentation_VISUAL.pdf

Speaker/reference text for each slide is read from the
<script type="text/x-notes"> ... </script> block inside that slide's HTML and
written into the PPTX slide notes, so the deck stays searchable and editable.
"""

import html
import re
import subprocess
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent          # ppt/visual
PPT = HERE.parent                                # ppt
EDGE = Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe")
USER_DATA = HERE / ".edge-profile"
N_SLIDES = 7
W, H = 2560, 1440                                # 2x of 1280x720
MIN_BYTES = 50 * 1024

PPTX_OUT = PPT / "archive" / "SIH26120_Idea_Presentation_VISUAL.pptx"
PDF_OUT = PPT / "archive" / "SIH26120_Idea_Presentation_VISUAL.pdf"


def shoot(html_path: Path, png_path: Path) -> None:
    """Render one HTML page to a 2560x1440 PNG with headless Edge."""
    if png_path.exists():
        png_path.unlink()
    cmd = [
        str(EDGE),
        "--headless=new",
        "--disable-gpu",
        "--no-sandbox",
        "--hide-scrollbars",
        "--force-color-profile=srgb",
        f"--user-data-dir={USER_DATA}",
        f"--screenshot={png_path}",
        "--window-size=1280,720",
        "--force-device-scale-factor=2",
        "--virtual-time-budget=4000",
        html_path.as_uri(),
    ]
    subprocess.run(cmd, capture_output=True, timeout=180)


def verify(png_path: Path) -> str:
    """Fail loudly if a PNG is missing, too small, or the wrong size."""
    if not png_path.exists():
        raise SystemExit(f"FAIL {png_path.name}: not produced")
    size = png_path.stat().st_size
    if size < MIN_BYTES:
        raise SystemExit(f"FAIL {png_path.name}: only {size} bytes (<50KB)")
    with Image.open(png_path) as im:
        dims = im.size
    if dims != (W, H):
        raise SystemExit(f"FAIL {png_path.name}: {dims} != {(W, H)}")
    return f"OK   {png_path.name}  {dims[0]}x{dims[1]}  {size / 1024:.0f} KB"


NOTES_RE = re.compile(
    r'<script[^>]*type=["\']text/x-notes["\'][^>]*>(.*?)</script>',
    re.S | re.I,
)


def notes_of(html_path: Path) -> str:
    m = NOTES_RE.search(html_path.read_text(encoding="utf-8"))
    return html.unescape(m.group(1)).strip() if m else ""


def build_pptx(pngs, htmls) -> None:
    prs = Presentation()
    prs.slide_width = Emu(12192000)              # 13.333in — 16:9
    prs.slide_height = Emu(6858000)              # 7.5in
    blank = prs.slide_layouts[6]
    for png, src in zip(pngs, htmls):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(png), 0, 0, width=prs.slide_width, height=prs.slide_height
        )
        text = notes_of(src)
        if text:
            slide.notes_slide.notes_text_frame.text = text
    prs.save(PPTX_OUT)


def build_pdf(pngs) -> None:
    frames = [Image.open(p).convert("RGB") for p in pngs]
    frames[0].save(
        PDF_OUT, "PDF", save_all=True, append_images=frames[1:], resolution=192.0
    )
    for f in frames:
        f.close()


def main() -> None:
    only = sys.argv[1:]                          # e.g. `render.py 3 5` re-shoots those
    htmls = [HERE / f"slide{i}.html" for i in range(1, N_SLIDES + 1)]
    pngs = [HERE / f"slide{i}.png" for i in range(1, N_SLIDES + 1)]

    for i, (src, png) in enumerate(zip(htmls, pngs), start=1):
        if only and str(i) not in only:
            continue
        if not src.exists():
            raise SystemExit(f"FAIL: {src.name} missing")
        shoot(src, png)

    for png in pngs:
        print(verify(png))

    build_pptx(pngs, htmls)
    build_pdf(pngs)
    print(f"\nPPTX  {PPTX_OUT}  {PPTX_OUT.stat().st_size / 1024:.0f} KB")
    print(f"PDF   {PDF_OUT}  {PDF_OUT.stat().st_size / 1024:.0f} KB")


if __name__ == "__main__":
    main()
